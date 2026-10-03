"""
Text extractors for different file types.
Yeh module PDF, DOCX, TXT, aur MD files se text nikalta hai.
"""

from pathlib import Path
from docx import Document
from pypdf import PdfReader


# ---------- Custom Exceptions ----------

class ExtractionError(Exception):
    """Jab text extract na ho sake."""
    pass


# ---------- Extractors ----------

def extract_from_pdf(file_path: Path) -> str:
    """PDF file se text nikalta hai."""
    try:
        reader = PdfReader(str(file_path))
        text_parts = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(f"--- Page {page_num} ---\n{page_text}")
        return "\n\n".join(text_parts)
    except Exception as e:
        raise ExtractionError(f"PDF padhne mein error: {e}") from e


def extract_from_docx(file_path: Path) -> str:
    """DOCX file se text nikalta hai."""
    try:
        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Tables se bhi text nikaalte hain
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ExtractionError(f"DOCX padhne mein error: {e}") from e


def extract_from_text(file_path: Path) -> str:
    """TXT ya MD file se text nikalta hai."""
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Agar UTF-8 nahi chal raha, to latin-1 try karo
        try:
            return file_path.read_text(encoding="latin-1")
        except Exception as e:
            raise ExtractionError(f"Text file padhne mein error: {e}") from e
    except Exception as e:
        raise ExtractionError(f"File read error: {e}") from e


# ---------- Main Dispatcher ----------

def extract_text(file_path: Path) -> str:
    """
    File extension dekh kar sahi extractor use karta hai.
    
    Args:
        file_path: Path to the file.
    
    Returns:
        Extracted text as string.
    
    Raises:
        ExtractionError: Agar file type supported nahi ya extraction fail ho.
    """
    extension = file_path.suffix.lower()

    extractors = {
        ".pdf": extract_from_pdf,
        ".docx": extract_from_docx,
        ".txt": extract_from_text,
        ".md": extract_from_text,
    }

    if extension not in extractors:
        raise ExtractionError(
            f"File type '{extension}' supported nahi hai. "
            f"Allowed: {', '.join(extractors.keys())}"
        )

    return extractors[extension](file_path)